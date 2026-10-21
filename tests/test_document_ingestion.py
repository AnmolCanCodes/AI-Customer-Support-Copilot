import asyncio
import io
import unittest

from docx import Document as DocxDocument
from fastapi import HTTPException
from openpyxl import Workbook

from src.ingestion import load_documents_from_uploaded_files, split_documents
import api


class UploadFileStub:
    def __init__(self, name, content):
        self.name = name
        self._content = content

    def getvalue(self):
        return self._content

    async def seek(self, offset, whence=0):
        return None

    async def read(self):
        return self._content


def make_xlsx_bytes(sheets):
    """sheets: {sheet_name: [row, row, ...]}"""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_docx_bytes(paragraphs, table_rows=None):
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = str(value)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class IngestionUploadTests(unittest.TestCase):
    def test_load_documents_from_uploaded_files_handles_text_files(self):
        uploaded = [UploadFileStub("policy.txt", b"Refunds are allowed within 14 days.")]

        documents = load_documents_from_uploaded_files(uploaded)

        self.assertEqual(len(documents), 1)
        self.assertIn("Refunds are allowed within 14 days.", documents[0].page_content)
        self.assertEqual(documents[0].metadata["source"], "policy.txt")

    def test_csv_rows_keep_column_names(self):
        csv_bytes = b"Product,Price,Stock\nPhone X,19999,12\nCase,499,\n"
        documents = load_documents_from_uploaded_files([UploadFileStub("items.csv", csv_bytes)])

        self.assertEqual(len(documents), 1)
        content = documents[0].page_content
        self.assertIn("Product: Phone X | Price: 19999 | Stock: 12", content)
        # Empty cell is skipped, not rendered as 'Stock: '
        self.assertIn("Product: Case | Price: 499", content)
        self.assertNotIn("Stock: \n", content)
        self.assertEqual(documents[0].metadata["source"], "items.csv")

    def test_csv_non_utf8_falls_back(self):
        csv_bytes = "Name,City\nJosé,Zürich\n".encode("cp1252")
        documents = load_documents_from_uploaded_files([UploadFileStub("people.csv", csv_bytes)])

        self.assertIn("Name: José | City: Zürich", documents[0].page_content)

    def test_xlsx_one_document_per_sheet_and_empty_sheet_skipped(self):
        xlsx_bytes = make_xlsx_bytes({
            "Prices": [["Product", "Price"], ["Phone X", 19999.0], ["Case", 499]],
            "Empty": [],
            "Stock": [["Product", "Units"], ["Phone X", 12]],
        })
        documents = load_documents_from_uploaded_files([UploadFileStub("data.xlsx", xlsx_bytes)])

        self.assertEqual([d.metadata["sheet"] for d in documents], ["Prices", "Stock"])
        self.assertIn("Product: Phone X | Price: 19999", documents[0].page_content)
        self.assertEqual(documents[0].metadata["source"], "data.xlsx")

    def test_docx_extracts_paragraphs_and_tables(self):
        docx_bytes = make_docx_bytes(
            ["Return policy", "Items can be returned within 14 days."],
            table_rows=[["Item", "Window"], ["Phone", "14 days"]],
        )
        documents = load_documents_from_uploaded_files([UploadFileStub("policy.docx", docx_bytes)])

        self.assertEqual(len(documents), 1)
        content = documents[0].page_content
        self.assertIn("Items can be returned within 14 days.", content)
        self.assertIn("Item: Phone | Window: 14 days", content)

    def test_unsupported_extension_is_skipped(self):
        documents = load_documents_from_uploaded_files([UploadFileStub("old.doc", b"binary")])
        self.assertEqual(documents, [])

    def test_corrupt_file_raises_clear_error(self):
        with self.assertRaises(ValueError) as ctx:
            load_documents_from_uploaded_files([UploadFileStub("bad.xlsx", b"not a real xlsx")])
        self.assertIn("bad.xlsx", str(ctx.exception))

    def test_split_uses_800_chunks(self):
        csv_rows = "\n".join(f"row{i},{'x' * 50}" for i in range(200))
        csv_bytes = f"id,value\n{csv_rows}\n".encode()
        documents = load_documents_from_uploaded_files([UploadFileStub("big.csv", csv_bytes)])
        chunks = split_documents(documents)

        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(c.page_content) <= 800 for c in chunks))


class ChatErrorHandlingTests(unittest.TestCase):
    def test_chat_raises_clear_hf_credit_error(self):
        async def trigger():
            api.vectorstore = object()
            original = api.create_rag_pipeline
            api.create_rag_pipeline = lambda _: (_ for _ in ()).throw(Exception("Client error '402 Payment Required' for url 'https://router.huggingface.co/v1/chat/completions'"))
            try:
                with self.assertRaises(HTTPException) as ctx:
                    await api.chat(api.ChatRequest(question="What is this?", session_id="s1"))
                self.assertEqual(ctx.exception.status_code, 402)
                self.assertIn("credits", str(ctx.exception.detail).lower())
            finally:
                api.create_rag_pipeline = original

        asyncio.run(trigger())