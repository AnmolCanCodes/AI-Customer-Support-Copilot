import asyncio
import csv
import inspect
import io
from pathlib import Path
from tempfile import NamedTemporaryFile

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter


BASE_DIR = Path(__file__).resolve().parent.parent
KNOWLEDGE_BASE_DIR = BASE_DIR / 'data' / 'knowledge_base'

MAX_ROWS_PER_SHEET = 10000


def _read_uploaded_bytes(uploaded_file):
    if hasattr(uploaded_file, 'content') and not callable(uploaded_file.content):
        return uploaded_file.content

    if hasattr(uploaded_file, 'getvalue'):
        return uploaded_file.getvalue()

    if hasattr(uploaded_file, 'read'):
        read_method = uploaded_file.read
        if inspect.iscoroutinefunction(read_method):
            if hasattr(uploaded_file, 'seek') and inspect.iscoroutinefunction(uploaded_file.seek):
                try:
                    asyncio.run(uploaded_file.seek(0))
                except RuntimeError:
                    loop = asyncio.new_event_loop()
                    loop.run_until_complete(uploaded_file.seek(0))
                    loop.close()
            try:
                return asyncio.run(read_method())
            except RuntimeError:
                loop = asyncio.new_event_loop()
                result = loop.run_until_complete(read_method())
                loop.close()
                return result

        if hasattr(uploaded_file, 'seek'):
            uploaded_file.seek(0)
        return uploaded_file.read()

    raise TypeError(f'Unsupported uploaded file object: {type(uploaded_file)}')


# Helpers
def _cell_to_text(value):
    """Convert a cell value to clean text. Returns '' for empty cells."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_to_text(rows):
    """
    Turn tabular rows into self-contained lines so every chunk keeps its
    column names, e.g. 'Product: Phone X | Price: 19999 | Stock: 12'.
    The first non-empty row is treated as the header.
    """
    header = None
    lines = []

    for row in rows:
        cells = [_cell_to_text(c) for c in row]
        if not any(cells):
            continue  # skip fully empty rows

        if header is None:
            header = [c or f'Column {i + 1}' for i, c in enumerate(cells)]
            continue

        parts = []
        for i, cell in enumerate(cells):
            if not cell:
                continue
            name = header[i] if i < len(header) else f'Column {i + 1}'
            parts.append(f'{name}: {cell}')
        if parts:
            lines.append(' | '.join(parts))

        if len(lines) >= MAX_ROWS_PER_SHEET:
            break

    return '\n'.join(lines)


def _decode_text(file_bytes):
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode('utf-8', errors='replace')


# Per-format loaders: (file_bytes, file_name) -> list[Document]

def _load_txt(file_bytes, file_name):
    text = file_bytes.decode('utf-8', errors='replace')
    return [Document(page_content=text, metadata={'source': file_name})]


def _load_pdf(file_bytes, file_name):
    with NamedTemporaryFile(suffix='.pdf', delete=False) as temp_file:
        temp_file.write(file_bytes)
        temp_path = Path(temp_file.name)

    try:
        loaded_docs = PyPDFLoader(str(temp_path)).load()
        for doc in loaded_docs:
            doc.metadata['source'] = file_name
        return loaded_docs
    finally:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)


def _load_csv(file_bytes, file_name):
    text = _decode_text(file_bytes)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel

    rows = csv.reader(io.StringIO(text, newline=''), dialect)
    content = _rows_to_text(rows)
    if not content:
        return []
    return [Document(page_content=content, metadata={'source': file_name})]


def _load_xlsx(file_bytes, file_name):
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    documents = []
    try:
        for sheet in workbook.worksheets:
            content = _rows_to_text(sheet.iter_rows(values_only=True))
            if not content:
                continue  # skip empty sheets
            documents.append(
                Document(
                    page_content=content,
                    metadata={'source': file_name, 'sheet': sheet.title},
                )
            )
    finally:
        workbook.close()
    return documents


def _load_docx(file_bytes, file_name):
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    docx_file = DocxDocument(io.BytesIO(file_bytes))
    blocks = []

    # Walk the body in order so tables stay next to the text around them.
    for child in docx_file.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            text = Paragraph(child, docx_file).text.strip()
            if text:
                blocks.append(text)
        elif tag == 'tbl':
            table = Table(child, docx_file)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            table_text = _rows_to_text(rows)
            if table_text:
                blocks.append(table_text)

    if not blocks:
        return []
    return [Document(page_content='\n\n'.join(blocks), metadata={'source': file_name})]


# To support a new format later: write a loader above and add it here.
LOADERS = {
    '.txt': _load_txt,
    '.pdf': _load_pdf,
    '.csv': _load_csv,
    '.xlsx': _load_xlsx,
    '.docx': _load_docx,
}

SUPPORTED_EXTENSIONS = set(LOADERS)



def load_documents_from_uploaded_files(uploaded_files):
    documents = []

    for uploaded_file in uploaded_files or []:
        file_name = getattr(uploaded_file, 'filename', getattr(uploaded_file, 'name', 'uploaded_file'))
        file_suffix = Path(file_name).suffix.lower()
        loader = LOADERS.get(file_suffix)
        if loader is None:
            continue

        file_bytes = _read_uploaded_bytes(uploaded_file)
        try:
            documents.extend(loader(file_bytes, file_name))
        except Exception as exc:
            raise ValueError(f'Could not read "{file_name}": {exc}') from exc

    return documents


def load_documents():
    documents = []
    for file_path in KNOWLEDGE_BASE_DIR.glob('*.txt'):
        loader = TextLoader(
            str(file_path),
            encoding='utf-8'
        )

        documents.extend(loader.load())
    return documents


def split_documents(documents):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=100
    )

    chunks = splitter.split_documents(documents)

    return chunks